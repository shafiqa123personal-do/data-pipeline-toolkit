"""Save lists of dicts to a nicely formatted Excel file, and read them back."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_COLUMN_WIDTH = 60


def save_excel(sheets, path):
    """sheets = {"Sheet name": [{column: value, ...}, ...]}. Returns the saved Path."""
    if not sheets:
        raise ValueError("save_excel needs at least one sheet")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31])      # Excel limit: 31 characters
        headers = _headers(rows)
        if not headers:
            headers, rows = ["Info"], [{"Info": "No data"}]
        ws.append(headers)
        for row in rows:
            ws.append([_clean(row.get(header)) for header in headers])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"           # scraped text, never a formula
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.auto_filter.ref = ws.dimensions         # filter buttons on the header row
        ws.freeze_panes = "A2"                     # header stays visible when scrolling
        _fit_columns(ws)

    try:
        wb.save(path)
    except PermissionError:
        raise PermissionError(
            f"Cannot save {path.name} - it is probably open in Excel. Close it and run again."
        ) from None
    return path


def read_sheet(path, sheet):
    """Read one sheet back as a list of dicts. Returns [] if the sheet is missing."""
    wb = load_workbook(path, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            return []
        rows = wb[sheet].iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            return []
        return [dict(zip(headers, row)) for row in rows]
    finally:
        wb.close()


def _headers(rows):
    headers = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    return headers


def _clean(value):
    """Remove control characters Excel cannot store."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _fit_columns(ws):
    for index, column in enumerate(ws.iter_cols(), start=1):
        longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        ws.column_dimensions[get_column_letter(index)].width = min(longest + 2, MAX_COLUMN_WIDTH)
