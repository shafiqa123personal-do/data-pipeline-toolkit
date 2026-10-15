import pytest
from openpyxl import load_workbook

from core.excel import read_sheet, save_excel


def test_save_excel_formats_sheets(tmp_path):
    path = save_excel(
        {"First": [{"Name": "A", "Price": 1.5}, {"Name": "B", "Price": 2.0}],
         "Second": [{"Metric": "Total", "Value": 2}]},
        tmp_path / "out" / "report.xlsx",
    )
    wb = load_workbook(path)
    assert wb.sheetnames == ["First", "Second"]
    ws = wb["First"]
    assert [c.value for c in ws[1]] == ["Name", "Price"]
    assert ws["A1"].font.bold is True
    assert ws.auto_filter.ref == "A1:B3"
    assert ws.freeze_panes == "A2"
    assert ws["B3"].value == 2.0
    assert ws.column_dimensions["A"].width >= 4


def test_save_excel_union_of_keys(tmp_path):
    path = save_excel({"S": [{"A": 1}, {"A": 2, "B": 3}]}, tmp_path / "r.xlsx")
    ws = load_workbook(path)["S"]
    assert [c.value for c in ws[1]] == ["A", "B"]
    assert ws["B3"].value == 3


def test_save_excel_empty_sheet(tmp_path):
    path = save_excel({"Empty": []}, tmp_path / "r.xlsx")
    ws = load_workbook(path)["Empty"]
    assert ws["A1"].value == "Info"
    assert ws["A2"].value == "No data"


def test_save_excel_no_sheets(tmp_path):
    with pytest.raises(ValueError):
        save_excel({}, tmp_path / "r.xlsx")


def test_save_excel_locked_file_message(tmp_path, monkeypatch):
    def locked(self, filename):
        raise PermissionError("locked")

    monkeypatch.setattr("core.excel.Workbook.save", locked)
    with pytest.raises(PermissionError, match="Close it and run again"):
        save_excel({"S": [{"A": 1}]}, tmp_path / "r.xlsx")


def test_read_sheet_round_trip(tmp_path):
    rows = [{"Title": "A", "Price": 10.0}, {"Title": "B", "Price": 20.5}]
    path = save_excel({"Books": rows}, tmp_path / "r.xlsx")
    assert read_sheet(path, "Books") == rows


def test_read_sheet_missing_sheet(tmp_path):
    path = save_excel({"Books": [{"A": 1}]}, tmp_path / "r.xlsx")
    assert read_sheet(path, "Nope") == []


def test_save_excel_formula_text_stays_text(tmp_path):
    path = save_excel({"S": [{"Name": "=HYPERLINK(\"http://evil\")"}]}, tmp_path / "r.xlsx")
    cell = load_workbook(path)["S"]["A2"]
    assert cell.data_type == "s"
    assert cell.value == "=HYPERLINK(\"http://evil\")"


def test_save_excel_strips_illegal_characters(tmp_path):
    path = save_excel({"S": [{"Name": "Bad\x01Name"}]}, tmp_path / "r.xlsx")
    assert load_workbook(path)["S"]["A2"].value == "BadName"
