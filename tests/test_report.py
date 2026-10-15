from datetime import datetime

import pytest
from openpyxl import load_workbook

from automation.report import build_report, compare_books, summarize_books
from core.excel import read_sheet


def book(title, price, rating=3, category="Poetry"):
    return {"Title": title, "Price": price, "Rating": rating, "Category": category,
            "Link": f"https://books.test/{title}"}


def as_dict(summary):
    return {row["Metric"]: row["Value"] for row in summary}


def test_summarize_books():
    summary = as_dict(summarize_books([book("A", 10.0, rating=5), book("B", 20.0, rating=3)]))
    assert summary["Total books"] == 2
    assert summary["Average price (£)"] == 15.0
    assert summary["Cheapest (£)"] == 10.0
    assert summary["Most expensive (£)"] == 20.0
    assert summary["5-star books"] == 1
    assert summary["3-star books"] == 1
    assert summary["1-star books"] == 0
    assert summary["Category: Poetry"] == 2


def test_summarize_books_empty():
    summary = as_dict(summarize_books([]))
    assert summary["Total books"] == 0
    assert summary["Average price (£)"] == ""


def test_compare_books_first_run():
    assert compare_books([book("A", 10.0)], None)[0]["Change"] == "No previous run"


def test_compare_books_detects_changes():
    previous = [book("A", 10.0), book("B", 20.0)]
    current = [book("A", 12.0), book("C", 5.0)]
    changes = {row["Change"]: row for row in compare_books(current, previous)}
    assert changes["Price changed"]["Title"] == "A"
    assert changes["Price changed"]["Old price"] == 10.0
    assert changes["Price changed"]["New price"] == 12.0
    assert changes["New"]["Title"] == "C"
    assert changes["Removed"]["Title"] == "B"


def test_compare_books_no_changes():
    rows = [book("A", 10.0)]
    assert compare_books(rows, rows)[0]["Change"] == "No changes"


def on_page(row, page):
    return {**row, "Page": page}


def test_compare_books_ignores_pages_missing_this_run():
    previous = [on_page(book("A", 10.0), 1), on_page(book("B", 20.0), 2)]
    current = [on_page(book("A", 10.0), 1)]
    assert compare_books(current, previous)[0]["Change"] == "No changes"


def test_compare_books_extra_pages_not_reported_as_new():
    previous = [on_page(book("A", 10.0), 1)]
    current = [on_page(book("A", 10.0), 1), on_page(book("C", 5.0), 2)]
    assert compare_books(current, previous)[0]["Change"] == "No changes"


def test_compare_books_same_pages_still_detects_removed():
    previous = [on_page(book("A", 10.0), 1), on_page(book("B", 20.0), 1)]
    current = [on_page(book("A", 10.0), 1)]
    assert compare_books(current, previous)[0]["Change"] == "Removed"


def test_build_report_first_run(tmp_path):
    path = build_report("books", [book("A", 10.0)], tmp_path, now=datetime(2026, 10, 15, 9, 30))
    assert path.name == "books_report_2026-10-15_0930.xlsx"
    assert load_workbook(path).sheetnames == ["Summary", "Books", "Changes"]
    assert read_sheet(path, "Changes")[0]["Change"] == "No previous run"


def test_build_report_second_run_tracks_changes(tmp_path):
    build_report("books", [book("A", 10.0)], tmp_path, now=datetime(2026, 10, 15, 9, 30))
    path = build_report("books", [book("A", 11.0)], tmp_path, now=datetime(2026, 10, 15, 10, 0))
    changes = read_sheet(path, "Changes")
    assert changes[0]["Change"] == "Price changed"
    assert changes[0]["Old price"] == 10.0


def test_build_report_ignores_corrupt_previous(tmp_path):
    (tmp_path / "books_report_2026-10-14_0900.xlsx").write_bytes(b"not an excel file")
    path = build_report("books", [book("A", 10.0)], tmp_path, now=datetime(2026, 10, 15, 9, 30))
    assert read_sheet(path, "Changes")[0]["Change"] == "No previous run"


def test_build_report_unknown_kind(tmp_path):
    with pytest.raises(ValueError):
        build_report("cars", [], tmp_path)


from automation.report import summarize_leads


def test_summarize_leads():
    leads = [{"Phone": "1", "Website": "w", "Email": "e"}, {"Phone": "2", "Website": "", "Email": ""},
             {"Phone": "", "Website": "", "Email": ""}]
    summary = as_dict(summarize_leads(leads))
    assert summary["Total leads"] == 3
    assert summary["With phone"] == "2 (67%)"
    assert summary["With website"] == "1 (33%)"
    assert summary["With email"] == "1 (33%)"


def test_summarize_leads_without_email_column():
    assert "With email" not in as_dict(summarize_leads([{"Phone": "1", "Website": ""}]))


def test_build_report_leads(tmp_path):
    path = build_report("leads", [{"Name": "A", "Phone": "1", "Website": ""}], tmp_path,
                        now=datetime(2026, 10, 15, 9, 30))
    assert path.name == "leads_report_2026-10-15_0930.xlsx"
    assert load_workbook(path).sheetnames == ["Summary", "Leads"]
